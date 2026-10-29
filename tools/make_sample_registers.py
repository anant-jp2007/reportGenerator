"""Build two sample production registers for trying the app: before (4 entries) and after (10).

    python tools/make_sample_registers.py [folder]   -> PRODUCTION before.xlsx, PRODUCTION after.xlsx
                                                        (default folder: samples/)

Same layout as the real register: sheets "Production 63/80/100", title row, headings in row 2,
month rows, TOTAL rows, and pre-typed report numbers on empty rows.
"""
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.chart import lookup, sdr_for  # noqa: E402
from app.production import parse_size  # noqa: E402

OUT = ROOT / "samples"
HEAD = ["Sr.\nNo.", "DATE", "SIZE", "REPORT NO.", "BATCH NUMBER", "Meters", "Total\nMeters", "Quantity\n( MT )", "Remark"]

# (grade, sr, date, size, batch, meters, remark) -- report no. is made from grade + sr
BEFORE = [
    ("PE 63", 1, date(2025, 4, 1), "32 X 6 ( SDR 13.6 )", "2025 04 01 001 1", 1400, ""),
    ("PE 63", 2, date(2025, 4, 4), "50 X 6 ( SDR 13.6 )", "2025 04 04 001 1", 300, ""),
    ("PE 80", 1, date(2025, 4, 1), "40 X 8 ( SDR 13.6 )", "2025 04 01 002 1", 1000, ""),
    ("PE 100", 1, date(2025, 4, 11), "50 X 6 ( SDR 21 )", "2025 04 11 002 1", 300, ""),
]
NEW = [
    # normal entry, last day of the month (testing end rolls into May)
    ("PE 63", 3, date(2025, 4, 30), "63 X 4 ( SDR 21 )", "2025 04 30 001 1", 600, ""),
    # PE 63 has no PN 10: can't be certified, listed with the reason
    ("PE 63", 4, date(2025, 5, 5), "40 X 10", "2025 05 05 001 1", 1600, ""),
    # SDR not written: taken from PE 80 + PN 10 = SDR 11
    ("PE 80", 2, date(2025, 5, 6), "40 X 10", "2025 05 06 002 1", 1000, ""),
    # PN with a decimal (12.5); testing end crosses into the next year
    ("PE 80", 3, date(2025, 12, 30), "63 X 12.5 ( SDR 9 )", "2025 12 30 001 1", 250, ""),
    # small "x", no SDR, no batch; last days of the financial year
    ("PE 100", 2, date(2026, 3, 29), "110 x 6", "", 400, "MARKET \nSAMPLE"),
    # new financial year: Sr. No. starts at 1 again; large size
    ("PE 100", 1, date(2026, 4, 2), "315 X 5 ( SDR 26 )", "2026 04 02 002 1", 120, ""),
]
EDITED_IN_AFTER = {("PE 63", 2): 350}     # meters corrected later: must NOT get a new certificate
PLACEHOLDERS = 3                          # empty rows with only Sr. No. and Report No. typed

thin = Side(style="thin")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = Font(name="Times New Roman", size=12)
BOLD = Font(name="Times New Roman", size=12, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONTH_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="92D050")


def tonnes(grade, size, meters):
    od, pn, sdr = parse_size(size)
    spec = lookup(grade, od, pn, sdr) if od and (sdr or sdr_for(grade, pn)) else None
    return round(meters * spec.weight / 1000, 3) if spec else None


def month_title(d):
    return d.strftime("%B - %Y")


def build(entries, path):
    wb = Workbook()
    wb.remove(wb.active)
    for grade in ("PE 63", "PE 80", "PE 100"):
        code = grade.split()[1]
        ws = wb.create_sheet(f"Production {code}")
        for col, w in zip("ABCDEFGHI", (6, 13, 22, 14, 20, 9, 9, 11, 12)):
            ws.column_dimensions[col].width = w

        ws.append([f"Production - {code}"])
        ws.merge_cells("A1:I1")
        ws["A1"].font, ws["A1"].alignment = Font(name="Times New Roman", size=16, bold=True), CENTER
        ws.append(HEAD)
        for c in ws[2]:
            c.font, c.alignment, c.border = BOLD, CENTER, BOX
        ws.row_dimensions[2].height = 32

        rows = sorted((e for e in entries if e[0] == grade), key=lambda e: e[2])
        month, month_m, month_t = None, 0, 0.0
        last_sr = 0

        def close_month():
            if month is None:
                return
            ws.append(["TOTAL ( MT )", None, None, None, None, None, month_m, round(month_t, 3)])
            r = ws.max_row
            ws.merge_cells(f"A{r}:F{r}")
            for c in ws[r]:
                c.font, c.alignment, c.border, c.fill = BOLD, CENTER, BOX, TOTAL_FILL

        for g, sr, d, size, batch, meters, remark in rows:
            if month_title(d) != month:
                close_month()
                month, month_m, month_t = month_title(d), 0, 0.0
                ws.append([month])
                r = ws.max_row
                ws.merge_cells(f"A{r}:I{r}")
                ws[f"A{r}"].font, ws[f"A{r}"].alignment, ws[f"A{r}"].fill = BOLD, CENTER, MONTH_FILL
            t = tonnes(g, size, meters)
            ws.append([sr, d, size, f"MP{code}/{sr:03d}", batch or None, meters, meters, t, remark or None])
            r = ws.max_row
            ws[f"B{r}"].number_format = "dd-mm-yyyy"
            ws[f"H{r}"].number_format = "0.000"
            for c in ws[r]:
                c.font, c.alignment, c.border = FONT, CENTER, BOX
            month_m += meters
            month_t += t or 0
            last_sr = sr
        close_month()

        for sr in range(last_sr + 1, last_sr + 1 + PLACEHOLDERS):
            ws.append([None, None, None, f"MP{code}/{sr:03d}"])
            for c in ws[ws.max_row]:
                c.font, c.alignment, c.border = FONT, CENTER, BOX

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(path)


def make(out_dir: Path = OUT) -> tuple[Path, Path]:
    """Write both registers into out_dir; returns (before, after)."""
    before, after_path = Path(out_dir) / "PRODUCTION before.xlsx", Path(out_dir) / "PRODUCTION after.xlsx"
    build(BEFORE, before)
    after = [(g, sr, d, s, b, EDITED_IN_AFTER.get((g, sr), m), rem) for g, sr, d, s, b, m, rem in BEFORE] + NEW
    build(after, after_path)
    return before, after_path


if __name__ == "__main__":
    make(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT)
