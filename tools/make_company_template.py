"""Build a company's certificate template from templates/test_report.xlsx.

    python tools/make_company_template.py abhay

The table (rows 6 onwards) stays exactly the same; only the header changes, as described by the
company's profile in config/settings.json:
- template_layout.merges: the header's merged cells (rows 1-5 are unmerged first);
- template_layout.fonts: {cell: [size, bold]} for the header cells (Times New Roman, centred, wrapped);
- template_layout.row_heights: {row: height in points}, e.g. a taller header and, to keep the page at
  100% on A4, shorter empty rows under the table;
- the header texts (and any fixed text the company replaces, '' = empty).
Logo and stamp are not stored in the template: every certificate places them (profile "images").
"""
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.render_excel import _rich  # noqa: E402
from app.settings import CELL_RE, _strip_help  # noqa: E402

import json  # noqa: E402

BASE = ROOT / "templates" / "test_report.xlsx"


def main(company_id: str) -> None:
    # read the settings without checking them: the template being built doesn't exist yet
    s = _strip_help(json.loads((ROOT / "config" / "settings.json").read_text(encoding="utf-8")))
    profile = s["company"]["profiles"][company_id]
    lay = profile["template_layout"]
    out = ROOT / profile["template"]
    if out.resolve() == BASE.resolve():
        sys.exit("This company uses the base template; build that one with tools/make_template.py.")

    wb = openpyxl.load_workbook(BASE, rich_text=True)   # rich_text keeps the superscript 2 / 3
    ws = wb.active
    for rng in [r for r in ws.merged_cells.ranges if r.max_row <= 5]:
        ws.unmerge_cells(str(rng))
    for row in ws.iter_rows(min_row=1, max_row=5):
        for c in row:
            c.value = None
    for rng in lay["merges"]:
        ws.merge_cells(rng)
    for ref, (size, bold) in lay["fonts"].items():
        cell = ws[ref]
        cell.font = Font(name="Times New Roman", size=size, bold=bold)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # header box as in the base: outline round rows 1-5, a line above the title (row 5); nothing inside
    # (unmerging leaves the old merged cells' inner edges behind)
    thin, none = Side(style="thin"), Side()
    for r in range(1, 6):
        for c in range(1, 13):
            ws.cell(r, c).border = Border(left=thin if c == 1 else none, right=thin if c == 12 else none,
                                          top=thin if r in (1, 5) else none, bottom=thin if r == 5 else none)
    for r, h in lay["row_heights"].items():
        ws.row_dimensions[int(r)].height = h

    for ref, text in profile.items():
        if CELL_RE.match(ref):
            ws[ref].value = _rich(text, copy(ws[ref].font)) if text else None

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"template written: {out}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
