"""Build templates/report2.xlsx (Report 2, the testing record) from a filled-in or blank sample.

    python tools/make_report2_template.py "Report All Test.xlsx"

The sample is an .xlsx (open an .xls in Excel and save it as .xlsx first). Its first sheet is used:
row labels in A:C, one entry per column in D:G. The entry columns are emptied and centred; the page
is set to A4 landscape, one page. Everything else (fonts, borders, row heights) is kept.
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "templates" / "report2.xlsx"
FIRST_ROW, LAST_ROW = 2, 38          # Designation of Pipe ... Signature of QCI
ENTRY_COLUMNS = "DEFG"


def main(sample: Path) -> None:
    wb = openpyxl.load_workbook(sample)
    for ws in wb.worksheets[1:]:
        wb.remove(ws)
    ws = wb.worksheets[0]
    ws.title = "Report 2"
    labels = {r: str(ws.cell(r, 2).value or "").strip() for r in (FIRST_ROW, LAST_ROW - 1)}
    if not labels[FIRST_ROW].startswith("Designation of Pipe") or labels[LAST_ROW - 1] != "Report Number":
        sys.exit(f"Not a Report 2 layout: B{FIRST_ROW}={labels[FIRST_ROW]!r}, B{LAST_ROW - 1}={labels[LAST_ROW - 1]!r}")
    for col in ENTRY_COLUMNS:
        for r in range(FIRST_ROW, LAST_ROW + 1):
            c = ws[f"{col}{r}"]
            c.value = None
            c.alignment = Alignment(horizontal="center", vertical="center")
    ws.print_area = f"A1:{ENTRY_COLUMNS[-1]}{LAST_ROW}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True
    wb.save(OUT)
    print(f"template written: {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
