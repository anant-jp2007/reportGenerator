"""Website: choose files -> check entries -> download, through Flask's test client.

    python tests/test_web.py [production.xls]
"""
import html as htmllib
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import xlrd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app import web  # noqa: E402

REAL = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/production_sample.xls"
TMP = Path(tempfile.mkdtemp(prefix="webtest_"))
web.OUTPUT = TMP / "output"
web.UPLOADS = TMP / "uploads"
client = web.app.test_client()


def text(resp):
    return htmllib.unescape(resp.get_data(as_text=True))

failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


def old_copy(blank: dict) -> bytes:
    """The real register with some rows reset to empty (only Sr. No. and Report No. left)."""
    src = xlrd.open_workbook(str(REAL))
    wb = openpyxl.Workbook(write_only=True)
    for sh in src.sheets():
        ws = wb.create_sheet(sh.name)
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = xlrd.xldate.xldate_as_datetime(cell.value, src.datemode) if cell.ctype == 3 else (cell.value or None)
                if (r + 1) in blank.get(sh.name, []) and c not in (0, 3):
                    v = None
                vals.append(v)
            ws.append(vals)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(old: bytes, new: bytes, old_name="old.xlsx", new_name="PRODUCTION.xls", company="balaji"):
    return client.post("/check", data={"old": (io.BytesIO(old), old_name), "new": (io.BytesIO(new), new_name),
                                       "company": company}, content_type="multipart/form-data")


real = REAL.read_bytes()

r = client.get("/")
html = text(r)
check("start page opens with both file boxes", r.status_code == 200 and "Old register" in html and "New register" in html)
check("app is not titled after a company", "<title>Test certificates</title>" in html and "BALAJI" not in html.split("<main>")[0])
check("company menu lists Balaji Polyplast, then Abhay Polyplast",
      '<option value="balaji"' in html and 0 < html.find(">Balaji Polyplast<") < html.find(">Abhay Polyplast<"))

r = client.post("/check", data={"new": (io.BytesIO(real), "PRODUCTION.xls"), "company": "balaji"},
                content_type="multipart/form-data")
check("missing old register: clear message", r.status_code == 400 and "Choose the old register" in text(r))

r = upload(real, real, old_name="old.xls", company="nobody")
check("unknown company: asks to choose one", r.status_code == 400 and "Choose the company" in text(r))

r = upload(b"hello", real, old_name="notes.txt")
check("not an Excel file: clear message", r.status_code == 400 and "isn't an Excel file" in text(r))

r = upload(b"not really excel", real, old_name="broken.xlsx")
check("broken Excel file: clear message", r.status_code == 400 and "couldn't be opened" in text(r))

# PE 63 Sr 171, PE 100 Sr 15 (110 x 6, no batch) and Sr 18, 19 are new; PE 63 row 343 (40 X 10) is new but can't be made
old = old_copy({"Production 63": [368, 343], "Production 100": [49, 52, 53]})
r = upload(old, real)
html = text(r)
check("review names the chosen company", "For <strong>Balaji Polyplast</strong>." in html)
check("review lists 4 new entries", r.status_code == 200 and "4 new entries to certify" in html, html[html.find("<h1"):][:80])
check("grouped by grade, PE 63 before PE 100", 0 < html.find(">PE 63<") < html.find(">PE 100<"))
check("Report No. pre-filled from the register's column D",
      all(f'value="{n}"' in html for n in ("MP63/171", "MP100/015", "MP100/018", "MP100/019")))
check("warnings shown (SDR taken from PN, empty batch)", "SDR not written" in html and "Batch number is empty" in html)
check("skipped entry shown with reason", "No certificate for 1 entry" in html and "PE 63 has no PN 10" in html)
token = html.split('name="token" value="')[1].split('"')[0]
ids = [s.split('"')[0] for s in html.split('name="row" value="')[1:]]

form = {"token": token, "row": ids, "files": "each"}
form.update({f"no-{ids[0]}": "BPP63/EX001", f"no-{ids[1]}": "7"})
r = client.post("/download", data=form)
check("download returns a ZIP", r.status_code == 200 and r.mimetype == "application/zip", r.status_code)
z = zipfile.ZipFile(io.BytesIO(r.data))
names = sorted(z.namelist())
check("one Excel file per certificate in a folder per grade, named from Report No. column",
      names == sorted(["PE 63/MP63-171.xlsx", "PE 100/MP100-015.xlsx", "PE 100/MP100-018.xlsx",
                       "PE 100/MP100-019.xlsx"]), names)
d7 = openpyxl.load_workbook(io.BytesIO(z.read("PE 63/MP63-171.xlsx"))).active["D7"].value
check("typed Report No. 'BPP63/EX001' printed on the certificate", d7 == "BPP63/EX001", d7)
d7 = openpyxl.load_workbook(io.BytesIO(z.read("PE 100/MP100-015.xlsx"))).active["D7"].value
check("typed Report No. 7 kept as a number", d7 == 7, d7)
d7 = openpyxl.load_workbook(io.BytesIO(z.read("PE 100/MP100-018.xlsx"))).active["D7"].value
check("untouched Report No. is the register's column D", d7 == "MP100/018", d7)
saved = Path(r.headers["X-Saved-To"])
check("a copy is saved in the output folder, same folders", r.headers["X-Certificates"] == "4" and
      sorted(p.relative_to(saved).as_posix() for p in saved.rglob("*.xlsx")) == names, r.headers.get("X-Saved-To"))

r = client.post("/download", data={"token": token, "row": ids[2:], "files": "one"})
z = zipfile.ZipFile(io.BytesIO(r.data))
check("only ticked entries (2 x PE 100), all in one file named after the grade", z.namelist() == ["PE 100.xlsx"] and
      len(openpyxl.load_workbook(io.BytesIO(z.read(z.namelist()[0]))).sheetnames) == 2, z.namelist())

r = client.post("/download", data={"token": token, "row": ids[:2], "files": "each", "format_choice": "1",
                                   "format": ["pdf", "word"]})
names = sorted(zipfile.ZipFile(io.BytesIO(r.data)).namelist())
check("PDF + Word ticked: a .pdf and a .docx per certificate, in grade folders",
      names == sorted(["PE 63/MP63-171.pdf", "PE 63/MP63-171.docx", "PE 100/MP100-015.pdf", "PE 100/MP100-015.docx"]),
      names)
r = client.post("/download", data={"token": token, "row": ids, "files": "one", "format_choice": "1",
                                   "format": ["pdf", "excel", "word"]})
names = sorted(zipfile.ZipFile(io.BytesIO(r.data)).namelist())
check("all in one file: one file per grade and format, only grades with entries",
      names == sorted(f"{g}.{e}" for g in ("PE 63", "PE 100") for e in ("docx", "pdf", "xlsx")), names)
r = client.post("/download", data={"token": token, "row": ids, "files": "each", "format_choice": "1"})
check("no format ticked: clear message", r.status_code == 400 and "Tick at least one format" in text(r))
html = text(upload(old, real))
check("check screen offers PDF (ticked), Excel and Word",
      'value="pdf" checked' in html and 'value="excel"' in html and 'value="word"' in html)

r = client.post("/download", data={"token": token, "files": "each"})
check("nothing ticked: clear message", r.status_code == 400 and "Tick at least one entry" in text(r))
r = client.post("/download", data={"token": "abc123", "row": ids, "files": "each"})
check("expired files: asks to choose again", r.status_code == 410 and "Choose the registers again" in text(r))

r = upload(real, real, old_name="PRODUCTION old.xls")
check("same file twice: 'No new entries' with a hint about swapped files",
      "No new entries" in text(r) and "swapped" in text(r))

# a second company: its own header on its certificates, nothing else changes
real_load = web.load_settings


def two_companies():
    s = real_load()
    demo = {"name": "Demo Pipes", "template": "templates/test_report.xlsx", "images": [], "A1": "DEMO PIPES", "A2": "(Test company)", "A3": "Plot 1", "A4": "Rajkot",
            "A5": "( TEST CERTIFICATE OF HDPE PIPES IS 4984 : 2016 CM/L - 0000000 )"}
    return {**s, "company": {**s["company"], "profiles": {**s["company"]["profiles"], "demo": demo}}}


web.load_settings = two_companies
html = text(client.get("/"))
check("second company appears in the menu, default still first",
      0 < html.find(">Balaji Polyplast<") < html.find(">Demo Pipes<"))
html = text(upload(old, real, company="demo"))
check("review names the second company", "For <strong>Demo Pipes</strong>." in html)
token = html.split('name="token" value="')[1].split('"')[0]
ids = [s.split('"')[0] for s in html.split('name="row" value="')[1:]]
z = zipfile.ZipFile(io.BytesIO(client.post("/download", data={"token": token, "row": ids[:1], "files": "each"}).data))
ws = openpyxl.load_workbook(io.BytesIO(z.read(z.namelist()[0]))).active
check("its certificate has its own header (A1, A5), the rest unchanged",
      ws["A1"].value == "DEMO PIPES" and "0000000" in ws["A5"].value and ws["C14"].value == "Colour", ws["A1"].value)
web.load_settings = real_load

# Abhay: its own header, logo and stamp in every format; files sorted by grade as for Balaji
html = text(upload(old, real, company="abhay"))
check("review names Abhay Polyplast", "For <strong>Abhay Polyplast</strong>." in html)
token = html.split('name="token" value="')[1].split('"')[0]
ids = [s.split('"')[0] for s in html.split('name="row" value="')[1:]]
z = zipfile.ZipFile(io.BytesIO(client.post("/download", data={"token": token, "row": ids, "files": "each", "format_choice": "1",
                                                              "format": ["pdf", "excel", "word"]}).data))
names = sorted(z.namelist())
ws = openpyxl.load_workbook(io.BytesIO(z.read("PE 63/MP63-171.xlsx"))).active
check("Abhay Excel: own header, no 'Tested by', logo + stamp, same table",
      ws["C1"].value == "ABHAY POLYPLAST" and ws["A1"].value is None and ws["G39"].value is None
      and len(ws._images) == 2 and ws["C14"].value == "Colour", (ws["C1"].value, len(ws._images)))
import pymupdf  # noqa: E402
page = pymupdf.open(stream=z.read("PE 63/MP63-171.pdf"), filetype="pdf")[0]
words = " ".join(page.get_text().split())
check("Abhay PDF: header texts, phone line, 2 pictures, no 'Tested by'",
      "ABHAY POLYPLAST" in words and "info@abhaypolyplast.com" in words and "Tested by" not in words
      and len(page.get_images()) == 2 and "MP63/171" in words, len(page.get_images()))
import docx  # noqa: E402
d = docx.Document(io.BytesIO(z.read("PE 63/MP63-171.docx")))
check("Abhay Word: header in the table, 2 floating pictures",
      "ABHAY POLYPLAST" in d.tables[0].rows[0].cells[2].text and d.element.xml.count("<wp:anchor") == 2)

print(f"\n(files in {TMP})")
print("ALL PASS" if not failures else f"{failures} FAILED")
sys.exit(1 if failures else 0)
