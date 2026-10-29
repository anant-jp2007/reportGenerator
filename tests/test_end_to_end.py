"""End to end: the website with the sample before/after registers, all three formats.

    python tests/test_end_to_end.py

The registers come from tools/make_sample_registers.py (4 entries before, 10 after). Of the 6 new
entries one can't be certified (40 X 10 on PE 63); an edited old entry and the rows with only a
Report No. must be ignored. Every certificate is opened and checked in Excel, PDF and Word.
"""
import html as htmllib
import io
import sys
import tempfile
import zipfile
from pathlib import Path

import docx
import openpyxl
import pymupdf

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))
from app import web  # noqa: E402
from make_sample_registers import make  # noqa: E402

TMP = Path(tempfile.mkdtemp(prefix="e2e_"))
web.OUTPUT, web.UPLOADS = TMP / "output", TMP / "uploads"
client = web.app.test_client()
failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


before, after = make(TMP / "registers")

# ---------- choose files -> check entries ----------
r = client.post("/check", data={"company": "balaji",
                                "old": (io.BytesIO(before.read_bytes()), before.name),
                                "new": (io.BytesIO(after.read_bytes()), after.name)},
                content_type="multipart/form-data")
page = htmllib.unescape(r.get_data(as_text=True))
check("check screen: 5 new entries to certify", "5 new entries to certify" in page)
check("the PE 63 '40 X 10' entry is listed as not certifiable, with the reason",
      "No certificate for 1 entry" in page and "PE 63 has no PN 10" in page)
check("edited old entry (PE 63 Sr. 2) is not offered", "Make certificate for Sr. 2\"" not in page
      or page.count("MP63-002") == 0)
check("rows with only a Report No. are not offered", "MP63-005" not in page and "MP80-004" not in page)
check("notes shown: SDR filled in, empty batch",
      "SDR not written; taken as SDR 11 from PE 80 PN 10" in page and "Batch number is empty" in page)
token = page.split('name="token" value="')[1].split('"')[0]
ids = [s.split('"')[0] for s in page.split('name="row" value="')[1:]]
check("entries in order PE 63, PE 80, PE 100", [i.split("-")[0] for i in ids] == ["PE63", "PE80", "PE80", "PE100", "PE100"], ids)

# what each certificate must say (file, report no., D8, D10, batch, grade, PN, DN, SDR)
EXPECT = {
    "MP63-003": ("1", "30-04-2025", "03-05-2025", "2025 04 30 001 1", "PE 63", "4", "63", "21"),
    "MP80-002": ("2", "06-05-2025", "09-05-2025", "2025 05 06 002 1", "PE 80", "10", "40", "11"),
    "MP80-003": ("3", "30-12-2025", "02-01-2026", "2025 12 30 001 1", "PE 80", "12.5", "63", "9"),
    "MP100-002": ("4", "29-03-2026", "01-04-2026", "", "PE 100", "6", "110", "21"),
    "MP100-001": ("5", "02-04-2026", "05-04-2026", "2026 04 02 002 1", "PE 100", "5", "315", "26"),
}
CELLS = ("D7", "D8", "D10", "D11", "J7", "J8", "J9", "J10")

# ---------- download: a file per certificate, all formats ----------
r = client.post("/download", data={"token": token, "row": ids, "files": "each", "format_choice": "1",
                                   "format": ["pdf", "excel", "word"]})
z = zipfile.ZipFile(io.BytesIO(r.data))
names = sorted(z.namelist())
check("15 files: PDF, Excel and Word for each of the 5 certificates",
      names == sorted(f"{n}.{e}" for n in EXPECT for e in ("pdf", "xlsx", "docx")), names)


def shown(c):
    v = c.value
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    if isinstance(v, (int, float)):
        fmt = c.number_format
        return f"{v:.{len(fmt.split('.')[1]) if '.' in fmt else 0}f}"
    return str(v)


mfr = {}
for name, want in EXPECT.items():
    ws = openpyxl.load_workbook(io.BytesIO(z.read(f"{name}.xlsx"))).active
    got = tuple(shown(ws[c]) for c in CELLS)
    check(f"{name}.xlsx header: report no., dates (+3 days), batch, grade, PN, DN, SDR", got == want, got)
    mfr[want[4]] = float(ws["G33"].value)

    pdf_text = " ".join(pymupdf.open(stream=z.read(f"{name}.pdf"), filetype="pdf")[0].get_text().split())
    missing = [w for w in want if w and w not in pdf_text]
    check(f"{name}.pdf shows the same header values, one page",
          not missing and len(pymupdf.open(stream=z.read(f"{name}.pdf"), filetype="pdf")) == 1, missing)

    d = docx.Document(io.BytesIO(z.read(f"{name}.docx")))
    word_text = " ".join(" ".join(c.text.split()) for row in d.tables[0].rows for c in row.cells)
    missing = [w for w in want if w and w not in word_text]
    check(f"{name}.docx shows the same header values", not missing, missing)

ws = openpyxl.load_workbook(io.BytesIO(z.read("MP100-001.xlsx"))).active
check("315 mm SDR 26: chart limits (pipe end 7.00, OD 315.00-316.90, wall 12.20-13.60, ovality 11.10)",
      [shown(ws[c]) for c in ("E20", "E24", "E25", "E26", "E27", "E28")]
      == ["7.00", "315.00", "316.90", "12.20", "13.60", "11.10"])
od = (ws["J24"].value, ws["J25"].value)
check("315 mm: OD obtained strictly inside 315.00-316.90, minimum below maximum", 315 < od[0] < od[1] < 316.9, od)
check("MFR in this run: PE 100 < PE 80 < PE 63", mfr["PE 100"] < mfr["PE 80"] < mfr["PE 63"], mfr)

# ---------- download: all in one file, each format ----------
r = client.post("/download", data={"token": token, "row": ids, "files": "one", "format_choice": "1",
                                   "format": ["pdf", "excel", "word"]})
z = zipfile.ZipFile(io.BytesIO(r.data))
by_ext = {n.rsplit(".", 1)[1]: n for n in z.namelist()}
check("one file per format", sorted(by_ext) == ["docx", "pdf", "xlsx"], z.namelist())
check("combined Excel: 5 sheets", len(openpyxl.load_workbook(io.BytesIO(z.read(by_ext["xlsx"]))).sheetnames) == 5)
check("combined PDF: 5 pages", len(pymupdf.open(stream=z.read(by_ext["pdf"]), filetype="pdf")) == 5)
check("combined Word: 5 certificates", len(docx.Document(io.BytesIO(z.read(by_ext["docx"]))).tables) == 5)

# ---------- unticking and typed report numbers ----------
r = client.post("/download", data={"token": token, "row": [ids[0], ids[4]], "files": "each", "format_choice": "1",
                                   "format": ["excel"], f"no-{ids[0]}": "BPP63/EX101", f"no-{ids[4]}": "77"})
z = zipfile.ZipFile(io.BytesIO(r.data))
got = {n: openpyxl.load_workbook(io.BytesIO(z.read(n))).active["D7"].value for n in z.namelist()}
check("only the 2 ticked entries, with the typed Report No.",
      got == {"MP63-003.xlsx": "BPP63/EX101", "MP100-001.xlsx": 77}, got)

print(f"\n(files in {TMP})")
print("ALL PASS" if not failures else f"{failures} FAILED")
sys.exit(1 if failures else 0)
