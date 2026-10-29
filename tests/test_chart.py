"""Chart lookups must reproduce the specification column of the issued sample reports.

    python tests/test_chart.py [test_report_sample.xlsx]

Each sheet of the sample (75 ... 315, all PE 100 / PN 5 / SDR 26) is compared cell by cell:
E20 pipe end, E24/E25 OD min/max, E26/E27 wall min/max, E28 ovality, J8 PN.
"""
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.chart import ChartError, lookup, sdr_for  # noqa: E402

SAMPLE = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/test_report_sample.xlsx"
failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


wb = openpyxl.load_workbook(SAMPLE, data_only=True)
for ws in wb.worksheets:
    grade, pn, od, sdr = ws["J7"].value, ws["J8"].value, ws["J9"].value, ws["J10"].value
    s = lookup(grade, od, pn, sdr)
    got = (s.pipe_end, s.od_min, s.od_max, s.wall_min, s.wall_max, s.ovality_max)
    want = tuple(float(ws[c].value) for c in ("E20", "E24", "E25", "E26", "E27", "E28"))
    check(f"sheet {ws.title}: spec column matches chart", got == want, f"chart {got} sample {want}")

# grade + PN -> SDR, including sizes written without SDR in the production sheet
check("PE 100 PN 6 -> SDR 21 (e.g. '110 x 6')", sdr_for("PE 100", 6) == 21)
check("PE 63 PN 8 -> SDR 11 (e.g. '63 X 8')", sdr_for("PE 63", 8) == 11)
check("PE 80 PN 10 -> SDR 11 (e.g. '40 X 10')", sdr_for("PE 80", 10) == 11)
check("PE 63 PN 4 -> SDR 21 (e.g. '90 X 4')", sdr_for("PE 63", 4) == 21)
check("PE 63 has no PN 10 ('40 X 10' on the PE 63 sheet)", sdr_for("PE 63", 10) is None)

s = lookup("PE 63", 32, 6, 13.6)
check("PE 63 32 X 6 (SDR 13.6): wall 2.4-2.8, OD 32-32.3", (s.wall_min, s.wall_max, s.od_min, s.od_max) == (2.4, 2.8, 32, 32.3))
check("PE 63 32 X 6 weight 0.228 kg/m (matches production MT)", round(s.weight, 3) == 0.228, s.weight)


def raises(fn, *a, **k):
    try:
        fn(*a, **k)
    except ChartError as e:
        return str(e)
    return None


msg = raises(lookup, "PE 100", 16, 5)
check("16 mm SDR 26 rejected (chart shows --)", msg and "not made" in msg, msg)
msg = raises(lookup, "PE 63", 32, 8, 13.6)
check("PN/SDR mismatch rejected", msg and "don't match" in msg, msg)
msg = raises(lookup, "PE 63", 40, 10)
check("PE 63 PN 10 rejected", msg and "no PN 10" in msg, msg)
msg = raises(lookup, "PE 100", 37, 5)
check("unknown OD rejected", msg and "not in IS 4984" in msg, msg)

print("\nALL PASS" if not failures else f"\n{failures} FAILED")
sys.exit(1 if failures else 0)
