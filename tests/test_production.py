"""Production reading and new-row detection, on the real register and on altered copies.

    python tests/test_production.py [production.xls]

The altered copies (.xlsx, written to a temp folder) have: a "Brand Name" column in the
middle, an extra title row above the heading, headings written differently, columns in a
different order. All must read exactly like the original.
"""
import sys
import tempfile
from pathlib import Path

import openpyxl
import xlrd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.production import find_new_rows, read_production  # noqa: E402

REAL = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/production_sample.xls"
TMP = Path(tempfile.mkdtemp(prefix="prodtest_"))
failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


# ---------- build altered copies ----------

def load_xls(path):
    wb = xlrd.open_workbook(str(path))
    out = {}
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            vals = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                elif v == "":
                    v = None
                vals.append(v)
            rows.append(vals)
        out[sh.name] = rows
    return out


ORIGINAL = load_xls(REAL)
HEADER_ROW = 1          # 0-based: row 2 holds the headings in every grade sheet
GRADE_SHEETS = ["Production 63", "Production 80", "Production 100"]


def copy_book(book):
    return {n: [list(r) for r in rows] for n, rows in book.items()}


def save(book, name):
    wb = openpyxl.Workbook(write_only=True)
    for sheet, rows in book.items():
        ws = wb.create_sheet(sheet)
        for r in rows:
            ws.append(r)
    p = TMP / name
    wb.save(p)
    return p


def blank(book, sheet, excel_rows):
    """Reset rows to how they look before data entry: only Sr. No. and Report No. typed."""
    for er in excel_rows:
        row = book[sheet][er - 1]
        for c in range(len(row)):
            if c not in (0, 3):
                row[c] = None
    return book


def insert_brand_column(book):
    for s in GRADE_SHEETS:
        for i, r in enumerate(book[s]):
            r.insert(3, "Brand Name" if i == HEADER_ROW else ("Balaji" if r and isinstance(r[0], float) else None))
    return book


def add_title_row(book):
    for s in GRADE_SHEETS:
        book[s].insert(0, ["BALAJI POLYPLAST - PRODUCTION REGISTER"])
    return book


def rename_headings(book, names=None):
    names = names or {"SR.NO": 0, "Date ": 1, "Size": 2, "Report Number": 3, "Batch No.": 4}
    for s in GRADE_SHEETS:
        hdr = next(r for r in book[s][:5] if r and str(r[0]).replace("\n", "").startswith("Sr"))
        for text, c in names.items():
            hdr[c] = text
    return book


def reorder_columns(book, order=(4, 2, 0, 1, 3, 5, 6, 7, 8)):
    """Batch, Size, Sr, Date, Report, ... (keeps every column, new positions)."""
    for s in GRADE_SHEETS:
        book[s] = [[(r + [None] * 9)[c] for c in order] + r[9:] for r in book[s]]
    return book


def signature(result):
    """Everything that matters about the rows read, independent of row position."""
    return sorted((r.grade, r.sr_no, r.date, r.size_text, r.report_no, r.batch,
                   r.spec.sdr if r.spec else None, tuple(r.problems), tuple(r.warnings)) for r in result.rows)


# ---------- 1. the real file ----------

real = read_production(REAL)
counts = {g: sum(r.grade == g for r in real.rows) for g in ("PE 63", "PE 80", "PE 100")}
check("real file: all filled rows read", counts == {"PE 63": 341, "PE 80": 144, "PE 100": 34}, counts)
check("real file: no sheet errors", not real.errors, real.errors)
bad = [r for r in real.rows if not r.ok]
check("real file: only PE 63 '40 X 10' is skipped",
      [(r.grade, r.excel_row, r.problems) for r in bad] == [("PE 63", 343, ["PE 63 has no PN 10 in IS 4984 chart"])],
      [(r.grade, r.excel_row, r.problems) for r in bad])
filled_sdr = [r for r in real.rows if any("SDR not written" in w for w in r.warnings)]
check("real file: SDR filled in for the 10 sizes written without it", len(filled_sdr) == 10,
      [(r.size_text, r.spec.sdr) for r in filled_sdr])
check("real file: '110 x 6' on PE 100 -> SDR 21",
      any(r.size_text == "110 x 6" and r.spec.sdr == 21 for r in filled_sdr))
no_batch = [(r.grade, r.sr_no) for r in real.rows if "batch number is empty" in r.warnings]
check("real file: 3 rows warned for empty batch", no_batch == [("PE 100", 13), ("PE 100", 14), ("PE 100", 15)], no_batch)

# ---------- 2. altered layouts read the same ----------

base = signature(real)
variants = {
    "Brand Name column in the middle": insert_brand_column(copy_book(ORIGINAL)),
    "title row above the heading": add_title_row(copy_book(ORIGINAL)),
    "headings written differently": rename_headings(copy_book(ORIGINAL)),
    "columns in another order": reorder_columns(copy_book(ORIGINAL)),
    "all of the above together": reorder_columns(
        rename_headings(add_title_row(insert_brand_column(copy_book(ORIGINAL))),
                        {"S.No.": 0, "DATE": 1, "SIZE": 2, "Report No": 4, "BATCH NO": 5})),
}
for label, book in variants.items():
    res = read_production(save(book, label.replace(" ", "_") + ".xlsx"))
    check(f"same result with {label}", signature(res) == base and not res.errors,
          res.errors or f"{len(res.rows)} rows")

# ---------- 3. headings that must stop a sheet, never guess ----------

book = copy_book(ORIGINAL)
for r in book["Production 80"]:
    if len(r) > 4:
        r[4] = None                                   # Batch Number column removed
res = read_production(save(book, "no_batch_col.xlsx"))
check("missing Batch column stops that sheet with a clear message",
      res.errors.get("PE 80") == "Production 80: 'Batch Number' column not found", res.errors)
check("...other sheets still read", {r.grade for r in res.rows} == {"PE 63", "PE 100"})

def set_heading(book, sheet, col, text):
    hdr = book[sheet][HEADER_ROW]
    hdr.extend([None] * (col + 1 - len(hdr)))
    hdr[col] = text


book = copy_book(ORIGINAL)
for s in GRADE_SHEETS:
    set_heading(book, s, 9, "Dispatch Date")          # an extra column at the end
res = read_production(save(book, "dispatch_date.xlsx"))
check("'Date' + 'Dispatch Date': exact 'Date' is used", signature(res) == base and not res.errors, res.errors)

book = copy_book(ORIGINAL)
for s in GRADE_SHEETS:
    set_heading(book, s, 1, "Date (Prod.)")
    set_heading(book, s, 9, "Dispatch Date")
res = read_production(save(book, "two_dates.xlsx"))
check("two unclear date columns stop the sheet instead of guessing",
      all("more than one column could be 'Date'" in res.errors.get(g, "") for g in ("PE 63", "PE 80", "PE 100")),
      res.errors.get("PE 63"))

# ---------- 4. new-row detection ----------

old_same = save(copy_book(ORIGINAL), "old_same.xlsx")
d = find_new_rows(old_same, REAL)
check("identical files give 0 new rows", not d.new and not d.errors, len(d.new))

old = save(blank(blank(copy_book(ORIGINAL), "Production 100", [52, 53]), "Production 63", [368]), "old_3blank.xlsx")
d = find_new_rows(old, REAL)
got = [(r.grade, r.sr_no, r.report_no) for r in d.new]
check("finds exactly the 3 newly filled rows, PE 63 first",
      got == [("PE 63", 171, "MP63/171"), ("PE 100", 18, "MP100/018"), ("PE 100", 19, "MP100/019")], got)

old = save(blank(copy_book(ORIGINAL), "Production 80", [66]), "old_fy.xlsx")
d = find_new_rows(old, REAL)
got = [(r.grade, r.sr_no, str(r.date)) for r in d.new]
check("2025 Sr 1 found even though 2024 also has Sr 1", got == [("PE 80", 1, "2025-04-01")], got)

book = copy_book(ORIGINAL)
for r in book["Production 100"]:
    if r and r[0] == 19.0:
        r[5] = 999.0                                    # meters edited on an old row
d = find_new_rows(old_same, save(book, "edited.xlsx"))
check("edited old row is ignored (no new report)", not d.new, len(d.new))

# new entries missing what the certificate needs: listed with the reason, never dropped silently
book = copy_book(ORIGINAL)
pe100 = book["Production 100"]
r52, r53 = pe100[51], pe100[52]          # Sr 18 and 19 (rows 52, 53)
r52[2] = None                            # size empty
r53[1] = "31/02/2025"                    # impossible date
old = save(blank(copy_book(ORIGINAL), "Production 100", [52, 53]), "old_missing.xlsx")
d = find_new_rows(old, save(book, "new_missing.xlsx"))
got = {r.sr_no: r.problems for r in d.new}
check("new entry with empty size is listed: 'size is empty'", got.get(18) == ["size is empty"], got)
bad = next(r for r in d.new if r.sr_no == 19)
check("new entry with unreadable date: still certified, note says how to write it, dates blank",
      bad.ok and bad.date is None and any("not readable" in w and "05-05-2025" in w for w in bad.warnings), bad.warnings)
book = copy_book(ORIGINAL)
book["Production 100"][52][1] = None     # date empty
book["Production 100"][51][2] = "63"     # PN missing from the size
d = find_new_rows(old, save(book, "new_missing2.xlsx"))
got = {r.sr_no: r.problems for r in d.new}
nodate = next(r for r in d.new if r.sr_no == 19)
check("new entry with empty date: still certified with a note", nodate.ok and any("date is empty" in w for w in nodate.warnings),
      nodate.warnings)
check("new entry with PN missing ('63') is listed: size not understood",
      got.get(18) and "not understood" in got[18][0], got)
check("rows with only a Sr. No. (placeholders) still ignored", all(r.sr_no in (18, 19) for r in d.new), [r.sr_no for r in d.new])

# the everyday case at scale: 10 rows added at the end and 31 earlier rows edited in the same file
book = copy_book(ORIGINAL)
pe63 = book["Production 63"]
filled = [i for i, r in enumerate(pe63) if r and isinstance(r[0], float) and len(r) > 2 and r[2]]
added = filled[-10:]                                  # the last 10 filled rows are the "new" ones
old = save(blank(copy_book(ORIGINAL), "Production 63", [i + 1 for i in added]), "old_310.xlsx")
for i in filled[19:50]:                               # rows 20-50 edited: meters, batch, remark
    r = pe63[i]
    r[5] = (r[5] or 0) + 100
    r[4] = (r[4] or "") + "9"
    r.extend([None] * (9 - len(r)))
    r[8] = "corrected"
d = find_new_rows(old, save(book, "new_310.xlsx"))
want = [int(pe63[i][0]) for i in added]
got = [r.sr_no for r in d.new]
check("10 new rows found, none of the 31 edited rows (old register 300+ rows)",
      got == want and len(d.new) == 10, (got, want))

new_other_layout = save(rename_headings(add_title_row(insert_brand_column(copy_book(ORIGINAL))),
                                        {"S.No.": 0, "DATE": 1, "SIZE": 2, "Report No": 4, "BATCH NO": 5}), "new_layout.xlsx")
old = save(blank(copy_book(ORIGINAL), "Production 100", [52, 53]), "old_2blank.xlsx")
d = find_new_rows(old, new_other_layout)
got = [(r.grade, r.sr_no) for r in d.new]
check("old and new files with different layouts still compare", got == [("PE 100", 18), ("PE 100", 19)], got)

book = copy_book(ORIGINAL)
for r in book["Production 80"]:
    if len(r) > 4:
        r[4] = None
d = find_new_rows(save(book, "old_broken80.xlsx"), REAL)
check("a sheet unreadable in the old file is skipped, not reported as all-new",
      all(r.grade != "PE 80" for r in d.new) and any("Production 80" in e for e in d.errors), d.errors)

# other companies' registers: sheets named "PE 63" etc., extra sheets, grades they don't make
def renamed(book, names):
    return {names.get(n, n): rows for n, rows in book.items()}


PE_NAMES = {"Production 63": "PE 63", "Production 80": "PE-80", "Production 100": "PE 100"}
old = save(renamed(blank(copy_book(ORIGINAL), "Production 100", [52, 53]), PE_NAMES) | {"Total": [["x"]]}, "old_pe.xlsx")
new = save(renamed(copy_book(ORIGINAL), PE_NAMES) | {"Total - Monthwise": [["x"]]}, "new_pe.xlsx")
d = find_new_rows(old, new)
got = [(r.grade, r.sr_no) for r in d.new]
check("sheets named 'PE 63' / 'PE-80' read as grades, other sheets ignored",
      got == [("PE 100", 18), ("PE 100", 19)] and not d.errors, (got, d.errors))

only100 = {"Production 100": ORIGINAL["Production 100"]}
d = find_new_rows(save(blank(copy_book(only100), "Production 100", [52, 53]), "old_100.xlsx"), save(copy_book(only100), "new_100.xlsx"))
got = [(r.grade, r.sr_no) for r in d.new]
check("company making only PE 100: other grades skipped silently", got == [("PE 100", 18), ("PE 100", 19)] and not d.errors,
      (got, d.errors))

d = find_new_rows(save(copy_book(only100), "old_only100.xlsx"), REAL)
check("a grade with a sheet in only one file is not compared (never all-new), with the reason",
      all(r.grade == "PE 100" for r in d.new) and any("old file: no sheet for PE 63" in e for e in d.errors), d.errors)

d = find_new_rows(save({"Sheet1": [["x"]]}, "old_none.xlsx"), save({"Sheet1": [["x"]]}, "new_none.xlsx"))
check("no grade sheets at all: error", d.new == [] and any("no PE 63" in e for e in d.errors), d.errors)

# Sr. No. is the first key. A register can repeat a Sr. No. (two rows typed with the same number,
# or last year's block copied below): each old row covers one new row, so the extra row is new.
book = copy_book(ORIGINAL)
pe100 = book["Production 100"]
row17 = next(i for i, r in enumerate(pe100) if r and r[0] == 17)
copy17 = list(pe100[row17]); copy17[4] = "BATCH-EXTRA"
old = save(copy_book(book) | {"Production 100": pe100[:row17 + 1] + [list(pe100[row17])]}, "old_dup.xlsx")
new = save(copy_book(book) | {"Production 100": pe100[:row17 + 1] + [list(pe100[row17]), copy17]}, "new_dup.xlsx")
d = find_new_rows(old, new)
got = [(r.grade, r.sr_no, r.batch) for r in d.new]
check("a repeated Sr. No. already in the old file twice is not new; a third, different row with it is",
      got == [("PE 100", 17, "BATCH-EXTRA")], got)
new = save(copy_book(book) | {"Production 100": pe100[:row17 + 1] + [copy17, list(pe100[row17])]}, "new_dup2.xlsx")
d = find_new_rows(old, new)
got = [(r.grade, r.sr_no, r.batch) for r in d.new]
check("the different row is the new one wherever it stands", got == [("PE 100", 17, "BATCH-EXTRA")], got)

# PN and SDR that don't match the chart: the PN is trusted (user's decision)
book = copy_book(ORIGINAL)
book["Production 63"][-1] = list(book["Production 63"][-1])
last = next(i for i in range(len(book["Production 63"]) - 1, 0, -1) if book["Production 63"][i] and book["Production 63"][i][2])
book["Production 63"][last][2] = "50 X 6 ( SDR 17 )"
row = next(r for r in read_production(save(book, "pn_clash.xlsx")).rows if r.grade == "PE 63" and r.excel_row == last + 1)
check("'50 X 6 ( SDR 17 )' on PE 63: certified as PN 6 / SDR 13.6, with a warning",
      row.ok and row.spec.sdr == 13.6 and row.spec.pn == 6 and any("PN trusted" in w for w in row.warnings),
      (row.problems, row.warnings))

print(f"\n(test copies in {TMP})")
print("ALL PASS" if not failures else f"{failures} FAILED")
sys.exit(1 if failures else 0)
