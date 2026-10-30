"""Excel output: every cell written back exactly as the report engine decided.

    python tests/test_excel.py [production.xls]
"""
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.production import read_production  # noqa: E402
from app.render_excel import write_excel, write_excel_combined  # noqa: E402
from app.report import build_reports  # noqa: E402

REAL = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/production_sample.xls"
TMP = Path(tempfile.mkdtemp(prefix="exceltest_"))
failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


def shown(xl):
    """What Excel displays for a cell, for the formats this app uses."""
    v = xl.value
    if v is None:
        return ""
    if isinstance(v, CellRichText):
        return "".join(b if isinstance(b, str) else
                       ("^{%s}" % b.text if b.font.vertAlign == "superscript" else b.text) for b in v)
    if isinstance(v, (datetime, date)):
        assert xl.number_format == "dd\\-mm\\-yyyy", xl.number_format
        return v.strftime("%d-%m-%Y")
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        fmt = xl.number_format
        d = len(fmt.split(".")[1]) if "." in fmt else 0
        return f"{v:.{d}f}"
    return str(v)


rows = read_production(REAL).rows
# a spread of real rows: each grade, small and large sizes, SDR written and filled in, empty batch
picked = [r for r in rows if r.ok][::25] + [r for r in rows if r.ok and (not r.batch or r.sdr is None)]
reports = build_reports(picked)
template = openpyxl.load_workbook(ROOT / "templates/test_report.xlsx", rich_text=True).active

mismatch, fmt_problems = [], []
for rep in reports:
    ws = openpyxl.load_workbook(write_excel(rep, TMP / f"{rep.file_name}.xlsx"), rich_text=True).active
    for ref, cell in rep.cells.items():
        if shown(ws[ref]) != cell.text:
            mismatch.append((rep.file_name, ref, shown(ws[ref]), cell.text))
    if ws.merged_cells.ranges != template.merged_cells.ranges:
        fmt_problems.append((rep.file_name, "merges differ"))
    if ws["J24"].font.name != "Times New Roman" or ws["A1"].font.sz != 20:
        fmt_problems.append((rep.file_name, "fonts lost"))
    if ws._images:
        fmt_problems.append((rep.file_name, "has a picture; 'Tested by' must stay blank to sign by hand"))
    if (str(ws.page_setup.paperSize), ws.page_setup.fitToWidth, ws.page_setup.fitToHeight) != ("9", 1, 1):  # 9 = A4
        fmt_problems.append((rep.file_name, "page setup", ws.page_setup.paperSize))

check(f"{len(reports)} Excel reports: every cell shows exactly the engine's text", not mismatch, mismatch[:5])
check("layout kept: merges, fonts, no signature image, A4 one page", not fmt_problems, fmt_problems[:5])

ws = openpyxl.load_workbook(TMP / f"{reports[0].file_name}.xlsx", rich_text=True).active
check("superscripts kept: Kgs/Cm², kg/m³", isinstance(ws["K8"].value, CellRichText)
      and shown(ws["K8"]) == "Kgs/Cm^{2}" and shown(ws["D35"]) == "940 To 960 kg/m^{3}")
check("numbers stored as numbers (Excel can calculate with them)",
      all(isinstance(ws[c].value, (int, float)) for c in ("J8", "J9", "J10", "J24", "G33", "G36")))
check("dates stored as dates, shown with dashes", isinstance(ws["D8"].value, datetime)
      and ws["D8"].number_format == "dd\\-mm\\-yyyy")
coil = reports[0].cells["G23"]
check("length '---', coiling as made, units kept",
      ws["G21"].value == "---" and shown(ws["G23"]) == coil.text and ws["J21"].value == "meters" and ws["J23"].value == "mm",
      (ws["G21"].value, shown(ws["G23"]), coil.text))

sdr136 = next(r for r in reports if r.cells["J10"].text == "13.6")
ws = openpyxl.load_workbook(TMP / f"{sdr136.file_name}.xlsx").active
check("SDR 13.6 shows as 13.6 in Excel (not 14)", shown(ws["J10"]) == "13.6", shown(ws["J10"]))

combined = openpyxl.load_workbook(write_excel_combined(reports, TMP / "Combined.xlsx"), rich_text=True)
check("combined workbook: one sheet per report, in order",
      len(combined.sheetnames) == len(reports) and combined.sheetnames[0] == reports[0].file_name, combined.sheetnames[:3])
bad = [(ws.title, ref) for ws, rep in zip(combined.worksheets, reports)
       for ref, cell in rep.cells.items() if shown(ws[ref]) != cell.text]
check("combined workbook: every sheet has its report's exact contents", not bad, bad[:5])
check("combined workbook: each sheet keeps merges and print area, no signature image",
      all(ws.merged_cells.ranges == template.merged_cells.ranges and not ws._images and ws.print_area
          for ws in combined.worksheets))

print(f"\n(files in {TMP})")
print("ALL PASS" if not failures else f"{failures} FAILED")
sys.exit(1 if failures else 0)
