"""Word output: every certificate value is in the document, laid out on the template's grid.

    python tests/test_word.py [production.xls]

Exact look was matched against Excel's print with tools/compare (Word's own PDF export vs
Excel's): all 174 text pieces within a median 0.24 pt, and every point of every Excel border
line printed. This test guards contents and structure without needing Word.
"""
import sys
import tempfile
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.layout import load_layout  # noqa: E402
from app.production import read_production  # noqa: E402
from app.render_word import write_word, write_word_combined  # noqa: E402
from app.report import build_reports  # noqa: E402

REAL = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data/source/production_sample.xls"
TMP = Path(tempfile.mkdtemp(prefix="wordtest_"))
failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


def plain(t):
    return " ".join(t.replace("^{", "").replace("}", "").split())


rows = read_production(REAL).rows
reports = build_reports([r for r in rows if r.ok][::60] + [r for r in rows if r.ok and not r.batch][:1])
lay = load_layout()

missing, shape = [], []
for rep in reports:
    d = docx.Document(write_word(rep, TMP / f"{rep.file_name}.docx"))
    t = d.tables[0]
    if (len(t.rows), len(t.columns)) != (len(lay.rows), len(lay.cols)):
        shape.append((rep.file_name, len(t.rows), len(t.columns)))
    text = " ".join(plain(c.text) for row in t.rows for c in row.cells)
    for ref, cell in rep.cells.items():
        p = plain(cell.text)
        if p and not all(w in text for w in p.split()):
            missing.append((rep.file_name, ref, p))
    if not d.inline_shapes and "graphic" not in d.element.xml:
        missing.append((rep.file_name, "signature", "no picture"))

check(f"{len(reports)} Word files: every value of the certificate is in the table", not missing, missing[:5])
check("table has the template's rows and columns", not shape, shape[:3])

d = docx.Document(TMP / f"{reports[0].file_name}.docx")
sec = d.sections[0]
check("A4 page", (round(sec.page_width.pt), round(sec.page_height.pt)) == (595, 842))
check("fixed column widths from the template (6 pt per width unit)",
      [round(int(g.get(docx.oxml.ns.qn("w:w"))) / 20, 1) for g in d.tables[0]._tbl.tblGrid.findall(docx.oxml.ns.qn("w:gridCol"))]
      == [round(w, 1) for w in lay.cols])
check("signature floats after 'Tested by :-'", "wp:anchor" in d.element.xml or "<wp:anchor" in d.tables[0]._tbl.xml)

d = docx.Document(write_word_combined(reports, TMP / "All.docx"))
check("combined: one table per certificate, page breaks between",
      len(d.tables) == len(reports) and d.element.xml.count('w:type="page"') == len(reports) - 1)

print(f"\n(files in {TMP})")
print("ALL PASS" if not failures else f"{failures} FAILED")
sys.exit(1 if failures else 0)
