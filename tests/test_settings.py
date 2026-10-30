"""Settings must load cleanly and describe exactly the template layout.

    python tests/test_settings.py
"""
import copy
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.layout import load_layout  # noqa: E402
from app.render_pdf import picture_position  # noqa: E402
from app.report import SYMBOLS, Picture  # noqa: E402
from app.settings import companies, company, company_profile, load_settings, validate  # noqa: E402

failures = 0


def check(name, cond, detail=""):
    global failures
    failures += not cond
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))


s = load_settings()
check("settings load with no problems", True)

def plain(v):
    if isinstance(v, CellRichText):
        return "".join(b if isinstance(b, str) else
                       "^{%s}" % b.text if b.font.vertAlign == "superscript" else
                       next((k for k, (f, ch) in SYMBOLS.items() if f == b.font.rFont and ch == b.text), b.text)
                       for b in v)
    return v


# Cells the app fills
filled = [s["header_cells"][k] for k in s["header_cells"]]
filled += [c for c in s["chart_cells"] if c != "decimals"]
for v in s["random_values"].values():
    filled += [v["cell"]] if "cell" in v else list(v["cells"].values())
filled += list(s["fixed_values"])
base = openpyxl.load_workbook(ROOT / "templates/test_report.xlsx", rich_text=True).active

for cid, name in companies(s):
    # Every text cell of the company's template is in settings (fixed texts + its header), same text.
    profile = company(s, cid)
    ws = openpyxl.load_workbook(ROOT / profile["template"], rich_text=True).active
    template_texts = {c.coordinate: plain(c.value) for row in ws.iter_rows(max_row=40) for c in row if c.value is not None}
    settings_texts = {k: v for k, v in {**s["fixed_text"], **company_profile(s, cid)}.items() if v != ""}
    missing = sorted(set(template_texts) - set(settings_texts))
    extra = sorted(set(settings_texts) - set(template_texts))
    diff = [c for c in template_texts if c in settings_texts and template_texts[c] != settings_texts[c]]
    check(f"{name}: every template text is in settings", not missing, missing)
    check(f"{name}: settings has no text the template lacks", not extra, extra)
    check(f"{name}: texts are identical", not diff, [(c, template_texts[c], settings_texts[c]) for c in diff])

    not_empty = [c for c in filled if ws[c].value is not None]
    check(f"{name}: all value cells are empty in the template", not not_empty, not_empty)
    check(f"{name}: PN/DN/SDR cells print decimals (13.6 not 14)",
          all(ws[c].number_format == "General" for c in ("J8", "J9", "J10")))

    # the table is the same for every company: rows 6-37 identical to the base template
    same = all(plain(ws.cell(r, c).value) == plain(base.cell(r, c).value) and all(repr(getattr(ws.cell(r, c), k)) == repr(getattr(base.cell(r, c), k))
                       for k in ("font", "border", "alignment", "number_format"))
               for r in range(6, 38) for c in range(1, 13)) \
        and all(ws.row_dimensions[r].height == base.row_dimensions[r].height for r in range(6, 38)) \
        and all(ws.column_dimensions[k].width == base.column_dimensions[k].width for k in "ABCDEFGHIJKL") \
        and {str(m) for m in ws.merged_cells.ranges if m.min_row >= 6} == {str(m) for m in base.merged_cells.ranges if m.min_row >= 6}
    check(f"{name}: table (rows 6-37) identical to the base template", same)

    # one A4 page at 100%: Excel only prints at full size when everything fits
    lay = load_layout((ROOT / profile["template"]).resolve())
    room = 842 - lay.page_margins[2] - lay.page_margins[3]
    check(f"{name}: page fits A4 at 100% ({lay.height:.1f} of {room:.1f} pt)", lay.height <= room - 1)
    pics_inside = all(picture_position(lay, Picture(Path(i["file"]), i["cell"], i["dx"], i["dy"], i["height"]))[1]
                      + i["height"] <= lay.height for i in profile.get("images", []))
    check(f"{name}: logo / stamp stay on the page", pics_inside)


# The checks catch broken edits.
def broken(edit):
    t = copy.deepcopy(s)
    edit(t)
    return validate(t)


check("overlapping MFR ranges are caught",
      any("overlap" in p for p in broken(lambda t: t["random_values"]["melt_flow_rate"]["by_grade"].__setitem__("PE 80", [0.2, 0.5]))))
check("min above max is caught",
      any("not valid" in p for p in broken(lambda t: t["random_values"]["density"].update(min=960, max=950))))
check("stripe width range with min above max is caught",
      any("not valid" in p for p in broken(lambda t: t["random_values"]["stripe_width"].__setitem__("range_above_limit", [5.5, 4.25]))))
check("two values in one cell is caught",
      any("already used" in p for p in broken(lambda t: t["random_values"]["density"].__setitem__("cell", "G33"))))
check("missing company profile is caught",
      any("not found" in p for p in broken(lambda t: t["company"].__setitem__("selected", "xyz"))))
check("company template missing is caught",
      any("template" in p for p in broken(lambda t: t["company"]["profiles"]["abhay"].__setitem__("template", "x.xlsx"))))
check("company header over a value cell is caught",
      any("already used" in p for p in broken(lambda t: t["company"]["profiles"]["abhay"].__setitem__("J24", "x"))))
check("missing logo file is caught",
      any("not found" in p for p in broken(lambda t: t["company"]["profiles"]["abhay"]["images"][0].__setitem__("file", "no.png"))))

print("\nALL PASS" if not failures else f"\n{failures} FAILED")
sys.exit(1 if failures else 0)
