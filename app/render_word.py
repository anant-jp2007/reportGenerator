"""Write reports as Word documents laid out like Excel's print of templates/test_report.xlsx.

    write_word(report, "MP100-019.docx")             # one certificate
    write_word_combined(reports, "All.docx")         # one page per certificate

The certificate is one Word table with the template's columns, rows, merged cells, borders and
fonts (from app/layout.py), so it can still be edited in Word.
"""
from __future__ import annotations

import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl.utils import get_column_letter

from app.layout import TABLE_SHIFT_X, TEMPLATE, Layout, load_layout
from reportlab.pdfbase import pdfmetrics

from app.render_pdf import (CELL_PAD, LINE_SPACING, SIGNATURE_HEIGHT, SIGNATURE_LIFT, SUPER_RISE, SUPER_SIZE,
                            SUPERSCRIPT, font_for)
from app.report import Report

ROOT = Path(__file__).resolve().parent.parent
BORDER_EIGHTHS = 8          # Word border width in 1/8 pt: 8 = 1 pt (Excel's PDF draws 0.96 pt)
WORD_LIFT = 1.1             # Word sets text about 1.1 pt lower in a line than Excel; lift it back (measured)

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "centerContinuous": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
_VALIGN = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
           "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}


def _set_cell_borders(cell, sides: dict[str, bool]) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if sides.get(side):
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(BORDER_EIGHTHS))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


class _RawCell:
    """Just enough of a cell for _set_cell_borders, around a row's own <w:tc>."""
    def __init__(self, tc):
        self._tc = tc


def _tc_at(row, grid_col: int):
    """The row's own <w:tc> covering grid column grid_col (python-docx's row.cells would give the
    top cell of a vertically merged area instead)."""
    col = 0
    for tc in row._tr.tc_lst:
        span = tc.grid_span
        if col <= grid_col < col + span:
            return tc
        col += span
    raise IndexError(grid_col)


def _fixed_layout(table, widths_pt) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(round(sum(widths_pt) * 20)))
    tblW.set(qn("w:type"), "dxa")
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", 0), ("bottom", 0), ("left", CELL_PAD), ("right", CELL_PAD)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(round(v * 20)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_pt):
        gc.set(qn("w:w"), str(round(w * 20)))


def _font(run, name: str, size: float, bold: bool, italic: bool) -> None:
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def _excel_spaces(line: str, continued: bool) -> tuple[str, int]:
    """(text, trailing spaces). After a line break Excel drops the leading space but aligns the
    line as if that space had moved to its end."""
    if continued and line.startswith(" "):
        line = line.lstrip() + " "
    stripped = line.rstrip(" ")
    return stripped, len(line) - len(stripped)


def _add_text_runs(para, line: str, style) -> None:
    """Runs for one line; ^{x} becomes a superscript drawn like Excel's (67% size, raised 41%)."""
    pos = 0
    for m in list(SUPERSCRIPT.finditer(line)) + [None]:
        end = m.start() if m else len(line)
        if end > pos:
            _font(para.add_run(line[pos:end]), style.font, style.size, style.bold, style.italic)
        if m:
            r = para.add_run(m.group(1))
            _font(r, style.font, round(style.size * SUPER_SIZE * 2) / 2, style.bold, style.italic)
            pos_el = OxmlElement("w:position")
            pos_el.set(qn("w:val"), str(round(style.size * SUPER_RISE * 2)))
            r._r.get_or_add_rPr().append(pos_el)
            pos = m.end()


def _write_text(cell, text: str, style, is_number: bool, box_width: float) -> None:
    """Fill a cell so Word places the text where Excel does.

    Two Excel habits Word doesn't share are reproduced with paragraph indents:
    - trailing spaces count when Excel centres or right-aligns ("      Maximum ");
    - text wider than its cell isn't wrapped (unless the cell wraps): it spills past the edge,
      to the right when left-aligned, both sides when centred ("Batch Number   :-").
    """
    if style.wrap:
        lines = [_excel_spaces(ln, i > 0) for i, ln in enumerate(text.split("\n"))]
    else:
        lines = [_excel_spaces(text.replace("\n", " "), False)]
    align = style.h_align if style.h_align not in (None, "general") else ("right" if is_number else "left")
    centred = align in ("center", "centerContinuous")
    font = font_for(style.font, style.bold, style.italic)
    room = box_width - 2 * CELL_PAD

    para = cell.paragraphs[0]
    for i, (line, trailing) in enumerate(lines):
        if i:
            para = cell.add_paragraph()
        pf = para.paragraph_format
        pf.alignment = _ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(style.size * LINE_SPACING)
        if i == len(lines) - 1:
            # lifts the block: a centred cell moves up by half of this, a bottom-aligned one by all of it
            pf.space_after = Pt(WORD_LIFT * (2 if style.v_align == "center" else 1))

        plain = SUPERSCRIPT.sub(lambda m: m.group(1), line)
        tail = pdfmetrics.stringWidth(" " * trailing, font, style.size) if (centred or align == "right") else 0
        over = pdfmetrics.stringWidth(plain, font, style.size) + tail - room + 0.5
        left = right = 0.0
        if over > 0 and not style.wrap:
            if centred:
                left = right = -over / 2
            elif align == "right":
                left = -over
            else:
                right = -over
        if tail:
            # move the text left by the spaces' width without narrowing the line
            if centred:
                left -= tail
            else:
                right += tail
                left -= tail
        if left:
            pf.left_indent = Pt(left)
        if right:
            pf.right_indent = Pt(right)
        _add_text_runs(para, line, style)



def _float_signature(run, x_pt: float, y_pt: float) -> None:
    """Turn the inline picture in `run` into one floating over the page at (x, y) from the table cell."""
    inline = run._r.find(".//" + qn("wp:inline"))
    anchor = OxmlElement("wp:anchor")
    for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"),
                 ("relativeHeight", "2"), ("behindDoc", "0"), ("locked", "0"), ("layoutInCell", "1"),
                 ("allowOverlap", "1")):
        anchor.set(k, v)
    sp = OxmlElement("wp:simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    anchor.append(sp)
    for tag, rel, off in (("wp:positionH", "column", x_pt), ("wp:positionV", "paragraph", y_pt)):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", rel)
        o = OxmlElement("wp:posOffset")
        o.text = str(round(off * 12700))
        pos.append(o)
        anchor.append(pos)
    anchor.append(copy.deepcopy(inline.find(qn("wp:extent"))))
    eff = OxmlElement("wp:effectExtent")
    for k in ("l", "t", "r", "b"):
        eff.set(k, "0")
    anchor.append(eff)
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        el = inline.find(qn(tag))
        if el is not None:
            anchor.append(copy.deepcopy(el))
    inline.getparent().replace(inline, anchor)


def _add_certificate(doc, report: Report, lay: Layout) -> None:
    import openpyxl
    ws = openpyxl.load_workbook(TEMPLATE).active
    nrows, ncols = len(lay.rows), len(lay.cols)
    table = doc.add_table(rows=nrows, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    _fixed_layout(table, lay.cols)
    left_m, right_m, _, _ = lay.page_margins
    # Word measures a table indent to the text of the first cell, so the cell padding is added back
    indent = ((595.3 - left_m - right_m - lay.width) / 2 + TABLE_SHIFT_X if lay.centered else 0) + CELL_PAD
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(round(indent * 20)))
    ind.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(ind)

    for r, h in enumerate(lay.rows):
        row = table.rows[r]
        row.height = Pt(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for c, w in enumerate(lay.cols):
            row.cells[c].width = Pt(w)

    # borders per template cell, outer edges of merged areas only (same rule as the PDF)
    areas = list(ws.merged_cells.ranges)
    for r in range(1, nrows + 1):
        for c in range(1, ncols + 1):
            b = ws.cell(r, c).border
            area = next((g for g in areas if g.min_row <= r <= g.max_row and g.min_col <= c <= g.max_col), None)
            _set_cell_borders(table.cell(r - 1, c - 1), {
                "top": bool(b.top.style) and (not area or r == area.min_row),
                "bottom": bool(b.bottom.style) and (not area or r == area.max_row),
                "left": bool(b.left.style) and (not area or c == area.min_col),
                "right": bool(b.right.style) and (not area or c == area.max_col),
            })
    # merge after borders so each merged cell keeps its outer edges
    for g in areas:
        if g.max_row <= nrows and g.max_col <= ncols:
            table.cell(g.min_row - 1, g.min_col - 1).merge(table.cell(g.max_row - 1, g.max_col - 1))
            # Word keeps one cell per row of the merged area, carrying only its first column's
            # borders: give each row the area's outer edges (right edge from the last column, etc.)
            for r in range(g.min_row, g.max_row + 1):
                first, last = ws.cell(r, g.min_col).border, ws.cell(r, g.max_col).border
                cols = [ws.cell(r, c).border for c in range(g.min_col, g.max_col + 1)]
                tc = _tc_at(table.rows[r - 1], g.min_col - 1)
                tcPr = tc.get_or_add_tcPr()
                old = tcPr.find(qn("w:tcBorders"))
                if old is not None:
                    tcPr.remove(old)
                _set_cell_borders(_RawCell(tc), {
                    "top": r == g.min_row and any(b.top.style for b in cols),
                    "bottom": r == g.max_row and any(b.bottom.style for b in cols),
                    "left": bool(first.left.style),
                    "right": bool(last.right.style),
                })

    for r in range(1, nrows + 1):
        for c in range(1, ncols + 1):
            ref = f"{get_column_letter(c)}{r}"
            if ref not in lay.boxes:
                continue        # inside a merged area: its first cell decides alignment and text
            cell = table.cell(r - 1, c - 1)
            st = lay.styles[ref]
            cell.vertical_alignment = _VALIGN.get(st.v_align, WD_CELL_VERTICAL_ALIGNMENT.BOTTOM)
            # empty cells still hold one paragraph: keep it tiny so it never pushes the row taller
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(min(st.size * LINE_SPACING, lay.rows[r - 1]))
            rc = report.cells.get(ref)
            if rc and rc.text and ref in lay.boxes:
                _write_text(cell, rc.text, st, isinstance(rc.value, (int, float)) and not isinstance(rc.value, bool),
                            lay.boxes[ref].w)

    sig = ROOT / report.signature_file
    if sig.is_file():
        from app.layout import merged_area_of
        min_col, min_row, max_col, _ = merged_area_of(report.signature_next_to)
        target = table.cell(min_row - 1, max_col)          # first cell after "Tested by :-"
        run = target.paragraphs[0].add_run()
        run.add_picture(str(sig), height=Pt(SIGNATURE_HEIGHT))
        _float_signature(run, 0, -SIGNATURE_LIFT)


def _new_document():
    lay = load_layout()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(595.3), Pt(841.9)
    left, right, top, bottom = lay.page_margins
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Pt(left), Pt(right), Pt(top), Pt(max(bottom, 0))
    sec.header_distance = sec.footer_distance = Pt(0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(1)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(1)
    return doc, lay


def _finish(doc):
    # Word always keeps a paragraph after a table: make it 1 pt so it never spills onto a new page
    last = doc.paragraphs[-1] if doc.paragraphs else doc.add_paragraph()
    last.paragraph_format.line_spacing = Pt(1)


def write_word(report: Report, path: str | Path) -> Path:
    doc, lay = _new_document()
    doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p) if doc.paragraphs else None
    _add_certificate(doc, report, lay)
    doc.add_paragraph()
    _finish(doc)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def write_word_combined(reports: list[Report], path: str | Path) -> Path:
    doc, lay = _new_document()
    for i, report in enumerate(reports):
        _add_certificate(doc, report, lay)
        p = doc.add_paragraph()
        if i < len(reports) - 1:
            p.add_run().add_break(WD_BREAK.PAGE)
    _finish(doc)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
