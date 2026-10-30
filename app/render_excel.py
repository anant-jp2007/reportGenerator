"""Write reports as Excel files, filling the company's template (templates/test_report*.xlsx).

    write_excel(report, "MP100-019.xlsx")              # one report, one file
    write_excel_combined(reports, "Reports.xlsx")       # one sheet per report

The template keeps all formatting (fonts, borders, merges, row heights, A4 page setup);
this only puts in the contents from app/report.py. Numbers are stored as real numbers
with a number format matching the printed decimals, dates as real dates shown dd-mm-yyyy.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import coordinate_to_tuple
from openpyxl.utils.units import pixels_to_EMU

from app.report import MARKUP, SYMBOLS, Report

PX_PER_PT = 4 / 3


def _rich(text: str, font) -> CellRichText | str:
    """'Kgs/Cm^{2}' -> rich text with a superscript 2 in the cell's own font; '✆' -> Wingdings ')'."""
    if not MARKUP.search(text):
        return text
    base = dict(rFont=font.name, sz=font.sz, b=font.b, i=font.i)
    parts, pos = [], 0
    for m in MARKUP.finditer(text):
        if m.start() > pos:
            parts.append(TextBlock(InlineFont(**base), text[pos:m.start()]))
        if m.group(1) is not None:
            parts.append(TextBlock(InlineFont(**base, vertAlign="superscript"), m.group(1)))
        else:
            sym_font, char = SYMBOLS[m.group(2)]
            parts.append(TextBlock(InlineFont(**{**base, "rFont": sym_font}), char))
        pos = m.end()
    if pos < len(text):
        parts.append(TextBlock(InlineFont(**base), text[pos:]))
    return CellRichText(parts)


def _number_format(text: str) -> str:
    if "." in text:
        return "0." + "0" * len(text.split(".")[1])
    return "0"


def _fill(ws, report: Report) -> None:
    for ref, cell in report.cells.items():
        xl = ws[ref]
        v = cell.value
        if v is None or cell.text == "":
            xl.value = None
        elif isinstance(v, date):
            xl.value = v
            xl.number_format = "dd\\-mm\\-yyyy"      # dashes on every PC, whatever Windows uses
        elif isinstance(v, bool) or not isinstance(v, (int, float)):
            xl.value = _rich(cell.text, xl.font)
        else:
            # format follows the printed text: 315.50 -> 0.00, 0.288 -> 0.000, 13.6 -> 0.0, 680 -> 0
            xl.value = v
            xl.number_format = _number_format(cell.text)
    _add_pictures(ws, report)


def _add_pictures(ws, report: Report) -> None:
    """The company's logo / stamp, top-left at its cell + (dx, dy) points."""
    for pic in report.pictures:
        img = Image(str(pic.file))
        h = round(pic.height * PX_PER_PT)
        img.width, img.height = round(img.width * h / img.height), h
        row, col = coordinate_to_tuple(pic.cell)
        marker = AnchorMarker(col=col - 1, colOff=pixels_to_EMU(round(pic.dx * PX_PER_PT)),
                              row=row - 1, rowOff=pixels_to_EMU(round(pic.dy * PX_PER_PT)))
        img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(img.width),
                                                                        pixels_to_EMU(img.height)))
        ws.add_image(img)


def write_excel(report: Report, path: str | Path) -> Path:
    wb = openpyxl.load_workbook(report.template, rich_text=True)
    ws = wb.active
    ws.title = _sheet_title(report, set())
    _fill(ws, report)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_excel_combined(reports: list[Report], path: str | Path) -> Path:
    """All reports in one workbook, one sheet each (sheet name = file name). One company per run."""
    wb = openpyxl.load_workbook(reports[0].template, rich_text=True)
    template_ws = wb.active
    used: set[str] = set()
    for report in reports:
        ws = wb.copy_worksheet(template_ws)
        ws.title = _sheet_title(report, used)
        ws.print_area = template_ws.print_area.split("!")[-1] if template_ws.print_area else "A1:L42"
        ws.page_setup.paperSize = template_ws.page_setup.paperSize
        ws.page_setup.orientation = template_ws.page_setup.orientation
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        _fill(ws, report)
    wb.remove(template_ws)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _sheet_title(report: Report, used: set[str]) -> str:
    """Excel sheet names: max 31 characters, no []:*?/\\ and unique in the workbook."""
    base = re.sub(r"[\[\]:*?/\\]", "-", report.file_name)[:31]
    title, n = base, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title, n = base[:31 - len(suffix)] + suffix, n + 1
    used.add(title.lower())
    return title
