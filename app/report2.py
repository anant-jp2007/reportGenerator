"""Report 2, the testing record: the same entries as Report 1, side by side, one per column.

    write_report2(reports, "Report 2.xlsx")                       # 10 entries -> sheets of 4, 4, 2
    write_report2(reports, "Report 2.xlsx", "last report 2.xls")  # fill its incomplete sheet first

Every value is taken from the Report 1 certificate built for that entry (app/report.py), so when both
reports are made they show exactly the same values. Which Report 1 cell goes in which row is set in
settings (report2.rows); rows Report 1 has no value for stay empty. All sheets go in one workbook.

An incomplete Report 2 (an earlier one with free columns in its last sheet) can be continued: its
filled columns are copied unchanged and the new entries go into the free columns, then new sheets.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

from app.render_excel import _number_format
from app.report import Report
from app.settings import ROOT, load_settings

DESIGNATION_LABEL = "Designation of Pipe"
LAST_LABEL = "Report Number"


class Report2Error(Exception):
    pass


@dataclass
class Sheet:
    title: str
    kept: list[dict[int, tuple]] = field(default_factory=list)    # columns copied from an earlier Report 2:
    #                                                                  row -> (value, number format)
    reports: list[Report] = field(default_factory=list)


def entry_values(report: Report, rows: dict[str, str]) -> dict[int, tuple[object, str]]:
    """row -> (value, printed text) for one entry, from its Report 1 cells."""
    out = {}
    for row, source in rows.items():
        if source == "grade_number":
            text = report.row.grade.replace("PE", "").strip()
            out[int(row)] = (int(text), text)
            continue
        cell = report.cells.get(source)
        if cell is None or cell.value is None or cell.text.strip() == "":
            continue                                    # e.g. ovality left blank for 710 mm and above
        out[int(row)] = (cell.value, cell.text.strip())
    return out


def _xls_cell(wb, sh, r: int, c: int):
    """(value, number format) of an .xls cell, or None when empty."""
    import xlrd
    cell = sh.cell(r, c)
    if cell.value == "":
        return None
    fmt = wb.format_map[wb.xf_list[sh.cell_xf_index(r, c)].format_key].format_str
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate.xldate_as_datetime(cell.value, wb.datemode), fmt
    return cell.value, fmt


def _grids(path: Path):
    """(sheet title, {(row, col): (value, number format)}) for every sheet; 1-based rows and columns."""
    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path), formatting_info=True)
        for sh in wb.sheets():
            yield sh.name, {(r + 1, c + 1): v for r in range(sh.nrows) for c in range(sh.ncols)
                            if (v := _xls_cell(wb, sh, r, c)) is not None}
    else:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        for ws in wb.worksheets:
            yield ws.title, {(c.row, c.column): (c.value, c.number_format) for row in ws.iter_rows() for c in row
                             if c.value not in (None, "")}


def read_incomplete(path: str | Path, settings: dict | None = None) -> Sheet:
    """The last Report 2 sheet of an earlier file, with its filled columns. Raises Report2Error when the
    file has no Report 2 sheet or its last one is already full."""
    s = settings or load_settings()
    conf = s["report2"]
    cols = [openpyxl.utils.column_index_from_string(c) for c in conf["entry_columns"]]
    rows = [int(r) for r in conf["rows"]]
    first, last = min(rows), max(rows)
    template = openpyxl.load_workbook(ROOT / conf["template"]).active
    label_row = {str(template.cell(r, 2).value or "").strip(): r for r in range(1, last + 2)}
    want_first, want_last = label_row_of(label_row, DESIGNATION_LABEL), label_row_of(label_row, LAST_LABEL)

    found = None
    try:
        for title, grid in _grids(Path(path)):
            labels = {r: str(v[0]).strip() for (r, c), v in grid.items() if c == 2}
            if labels.get(want_first, "").startswith(DESIGNATION_LABEL) and labels.get(want_last) == LAST_LABEL:
                found = (title, grid)
    except Exception as e:                      # not a spreadsheet, damaged, password...
        raise Report2Error(f"the incomplete Report 2 couldn't be opened ({e})")
    if found is None:
        raise Report2Error("the incomplete Report 2 has no sheet laid out like Report 2 "
                           f"('{DESIGNATION_LABEL}' in B{want_first}, '{LAST_LABEL}' in B{want_last})")
    title, grid = found
    kept = []
    for c in cols:
        column = {r: grid[(r, c)] for r in range(first, last + 2) if (r, c) in grid}
        if not column:
            break                               # first free column: the rest are taken as free too
        kept.append(column)
    if len(kept) == len(cols):
        raise Report2Error(f"the last sheet of the incomplete Report 2 ('{title}') is already full")
    return Sheet(title, kept)


def label_row_of(label_row: dict[str, int], label: str) -> int:
    return next(r for text, r in label_row.items() if text.startswith(label))


def plan_sheets(reports: list[Report], per_sheet: int, start: Sheet | None = None) -> list[Sheet]:
    """Fill the free columns of `start` first, then new sheets of `per_sheet` entries."""
    sheets, todo = [], list(reports)
    if start is not None:
        free = per_sheet - len(start.kept)
        start.reports, todo = todo[:free], todo[free:]
        sheets.append(start)
    while todo:
        sheets.append(Sheet("", reports=todo[:per_sheet]))
        todo = todo[per_sheet:]
    used = {sh.title.lower() for sh in sheets if sh.title}
    # continuing "Report 2 - 3": the next sheets are 4, 5... so names stay unique across files
    m = re.fullmatch(r"Report 2 - (\d+)", start.title) if start is not None else None
    n = int(m.group(1)) + 1 if m else 1
    for sh in sheets:
        while not sh.title:
            candidate = f"Report 2 - {n}"
            n += 1
            if candidate.lower() not in used:
                sh.title = candidate
                used.add(candidate.lower())
    return sheets


def _put(xl, value, text: str) -> None:
    if isinstance(value, (date, datetime)):
        xl.value = value
        xl.number_format = "dd\\-mm\\-yyyy"
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        xl.value = value
        xl.number_format = _number_format(text)
    else:
        xl.value = text


def write_report2(reports: list[Report], path: str | Path, incomplete: str | Path | None = None,
                  settings: dict | None = None) -> tuple[Path, list[Sheet]]:
    s = settings or load_settings()
    conf = s["report2"]
    cols = conf["entry_columns"]
    start = read_incomplete(incomplete, s) if incomplete else None
    sheets = plan_sheets(reports, len(cols), start)

    wb = openpyxl.load_workbook(ROOT / conf["template"])
    template_ws = wb.active
    for sh in sheets:
        ws = wb.copy_worksheet(template_ws)
        ws.title = sh.title[:31]
        ws.print_area = template_ws.print_area.split("!")[-1] if template_ws.print_area else None
        ws.page_setup.paperSize = template_ws.page_setup.paperSize
        ws.page_setup.orientation = template_ws.page_setup.orientation
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        ws.print_options.horizontalCentered = template_ws.print_options.horizontalCentered
        for col, column in zip(cols, sh.kept):
            for r, (v, fmt) in column.items():          # copied as they were, format included
                xl = ws[f"{col}{r}"]
                xl.value = v
                if fmt and fmt != "General":
                    xl.number_format = fmt
        for col, rep in zip(cols[len(sh.kept):], sh.reports):
            for r, (v, text) in entry_values(rep, conf["rows"]).items():
                _put(ws[f"{col}{r}"], v, text)
    wb.remove(template_ws)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path, sheets
