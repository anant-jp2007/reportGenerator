"""The certificate's layout, read from templates/test_report.xlsx, for the PDF and Word writers.

Excel is the reference: the PDF and Word certificates must look like Excel's own print of the
template. The numbers below were measured from Excel's PDF export of this template
(see docs/report-rules.md, "Matching Excel"):
- column width: 6 pt per width unit (Times New Roman 12 is the workbook font: 8 px per unit),
  rounded down to 0.5 pt like Excel;
- fonts print at their own size (the sheet fits A4 at 100%);
- rows print at 97.4% of their set height in Excel's PDF;
- thin borders are 0.96 pt (as drawn in its PDF);
- the table is centred between the 0.4" side margins, 0.2165" from the top.
"""
from __future__ import annotations

import math
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "templates" / "test_report.xlsx"

PT_PER_WIDTH_UNIT = 6.0
ROW_SCALE = 0.974
THIN_LINE = 0.96
TABLE_SHIFT_X = -0.66     # Excel's centring lands 0.66 pt left of the exact middle
DEFAULT_ROW_HEIGHT = 15.75


@dataclass(frozen=True)
class Style:
    font: str
    size: float
    bold: bool
    italic: bool
    h_align: str | None        # left / center / right / None (Excel "general")
    v_align: str | None        # top / center / bottom
    wrap: bool
    number_format: str


@dataclass(frozen=True)
class Box:
    """A cell or merged area: position in points from the table's top-left corner."""
    ref: str
    x: float
    y: float
    w: float
    h: float


@dataclass
class Layout:
    cols: list[float]                     # widths of A..last column, points
    rows: list[float]                     # heights of row 1..last row, points
    boxes: dict[str, Box]                 # every cell that can hold text (merged areas as one box)
    styles: dict[str, Style]
    borders: list[tuple[float, float, float, float]]   # line segments (x1, y1, x2, y2), points, y down
    page_margins: tuple[float, float, float, float]    # left, right, top, bottom in points
    centered: bool

    @property
    def width(self) -> float:
        return sum(self.cols)

    @property
    def height(self) -> float:
        return sum(self.rows)


def _col_widths(ws, ncols: int) -> list[float]:
    widths = {}
    for dim in ws.column_dimensions.values():
        if dim.width:
            for c in range(dim.min, dim.max + 1):
                widths[c] = dim.width
    base = ws.sheet_format.defaultColWidth or (ws.sheet_format.baseColWidth or 8) + 0.71
    return [math.floor(widths.get(c, base) * PT_PER_WIDTH_UNIT * 2) / 2 for c in range(1, ncols + 1)]


@lru_cache(maxsize=4)
def load_layout(path: Path = TEMPLATE) -> Layout:
    ws = openpyxl.load_workbook(path, rich_text=True).active
    min_col, min_row, max_col, max_row = range_boundaries(ws.print_area[0].split("!")[-1].replace("$", "")
                                                          if isinstance(ws.print_area, list) else
                                                          str(ws.print_area).split("!")[-1].replace("$", ""))
    cols = _col_widths(ws, max_col)
    rows = [(ws.row_dimensions[r].height or DEFAULT_ROW_HEIGHT) * ROW_SCALE for r in range(1, max_row + 1)]
    xs = [0.0]
    for w in cols:
        xs.append(xs[-1] + w)
    ys = [0.0]
    for h in rows:
        ys.append(ys[-1] + h)

    merged = {}
    covered = set()
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                covered.add((r, c))

    boxes, styles, borders = {}, {}, []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            ref = f"{get_column_letter(c)}{r}"
            rng = merged.get((r, c))
            if rng:
                boxes[ref] = Box(ref, xs[c - 1], ys[r - 1], xs[rng.max_col] - xs[c - 1], ys[rng.max_row] - ys[r - 1])
            elif (r, c) not in covered:
                boxes[ref] = Box(ref, xs[c - 1], ys[r - 1], cols[c - 1], rows[r - 1])
            f, a = cell.font, cell.alignment
            styles[ref] = Style(f.name or "Times New Roman", float(f.sz or 12), bool(f.b), bool(f.i),
                                a.horizontal, a.vertical or "bottom", bool(a.wrap_text), cell.number_format)

            # borders: a side inside a merged area is not drawn (Excel shows only the outer edge)
            area = next((g for g in ws.merged_cells.ranges if (r, c) in _cells(g)), None)
            b = cell.border
            x1, x2, y1, y2 = xs[c - 1], xs[c], ys[r - 1], ys[r]
            if b.top.style and (not area or r == area.min_row):
                borders.append((x1, y1, x2, y1))
            if b.bottom.style and (not area or r == area.max_row):
                borders.append((x1, y2, x2, y2))
            if b.left.style and (not area or c == area.min_col):
                borders.append((x1, y1, x1, y2))
            if b.right.style and (not area or c == area.max_col):
                borders.append((x2, y1, x2, y2))

    m = ws.page_margins
    return Layout(cols, rows, boxes, styles, sorted(set(borders)),
                  (m.left * 72, m.right * 72, m.top * 72, m.bottom * 72), bool(ws.print_options.horizontalCentered))


def _cells(rng):
    return {(r, c) for r in range(rng.min_row, rng.max_row + 1) for c in range(rng.min_col, rng.max_col + 1)}


def merged_area_of(ref: str, path: Path = TEMPLATE):
    """(min_col, min_row, max_col, max_row) of the merged area starting at ref, or of the cell itself."""
    ws = openpyxl.load_workbook(path).active
    for rng in ws.merged_cells.ranges:
        if rng.coord.split(":")[0] == ref:
            return rng.min_col, rng.min_row, rng.max_col, rng.max_row
    c, r, _, _ = range_boundaries(ref + ":" + ref)
    return c, r, c, r
